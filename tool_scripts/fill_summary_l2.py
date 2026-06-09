"""Fill outputs/KB-l2_AdaExplore_50/Summary.xlsx with one row per test in test_list_2.txt.

Columns: Name | trt_baseline (ms) | Agent result | ratio | Fail | Implementation
"""

import json
import os
import re
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
LEVEL_DIR = os.path.join(ROOT, "datasets", "KernelBench", "level2")
TRT_JSON = os.path.join(ROOT, "results", "timing", "RTX-4090-D", "baseline_time_trt.json")
TEST_LIST = os.path.join(ROOT, "config", "test_list", "test_list_2.txt")
OUT_DIR = os.path.join(ROOT, "outputs", "KB-l2_AdaExplore_50")
XLSX = os.path.join(OUT_DIR, "Summary.xlsx")


def build_pid_to_filename():
    mapping = {}
    for fn in os.listdir(LEVEL_DIR):
        if not fn.endswith(".py"):
            continue
        m = re.match(r"^(\d+)_", fn)
        if not m:
            continue
        mapping[int(m.group(1))] = fn
    return mapping


def detect_impl(kernel_path):
    if not os.path.isfile(kernel_path):
        return ""
    try:
        with open(kernel_path, "r") as f:
            src = f.read()
    except Exception:
        return ""
    has_triton = "import triton" in src or "triton.jit" in src
    has_cuda = (
        "load_inline" in src
        or "cpp_extension" in src
        or "extern \"C\"" in src
        or "__global__" in src
        or "from numba import cuda" in src
    )
    tags = []
    if has_cuda:
        tags.append("CUDA")
    if has_triton:
        tags.append("Triton")
    if not tags:
        tags.append("PyTorch")
    return "+".join(tags)


def trt_fail_reason(entry):
    if entry is None:
        return "TRT entry missing"
    status = entry.get("status", "")
    if status == "onnx_export_failed":
        return "ONNX conversion failure"
    if status == "trt_build_failed":
        return "TensorRT execution failure"
    if status == "numeric_mismatch":
        return ""
    return ""


def classify_agent_error(err_name, err_msg):
    if err_name == "torch.OutOfMemoryError":
        return "CUDA OOM"
    if err_name == "triton.compiler.errors.CompilationError":
        return "Triton compile error"
    if err_name == "subprocess.CalledProcessError":
        return "C/CUDA compile error"
    if err_name == "builtins.RuntimeError":
        msg = (err_msg or "").lower()
        if "out of memory" in msg:
            return "CUDA OOM"
        if "triton error" in msg:
            return "Triton CUDA error"
        if "cuda" in msg:
            return "CUDA runtime error"
        return "Runtime error"
    if err_name:
        return err_name.rsplit(".", 1)[-1]
    return ""


def agent_fail_reason(metrics):
    if metrics is None:
        return "Agent optimization failure"
    compiled = metrics.get("compiled", False)
    correct = metrics.get("correctness", False)
    if compiled and correct:
        return ""
    meta = metrics.get("metadata", {}) or {}
    err_name = meta.get("runtime_error_name", "") or ""
    err_msg = meta.get("runtime_error", "") or ""
    label = classify_agent_error(err_name, err_msg)
    if not compiled:
        return label or "Compile failed"
    return label or "Wrong output"


def main():
    sys.path.insert(0, "/home/agiuser/.local/lib/python3.10/site-packages")
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill

    pid2fn = build_pid_to_filename()
    with open(TRT_JSON, "r") as f:
        trt = json.load(f).get("level2", {})

    rows = []
    with open(TEST_LIST, "r") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            parts = line.split()
            if len(parts) != 2:
                continue
            level, pid = int(parts[0]), int(parts[1])
            if level != 2:
                continue
            fn = pid2fn.get(pid)
            if fn is None:
                rows.append((f"(missing pid {pid})", "", "", "", "Test file not found", ""))
                continue

            trt_entry = trt.get(fn)
            trt_mean = trt_entry.get("mean") if trt_entry else None
            trt_ok = trt_entry is not None and trt_entry.get("status") == "ok" and trt_mean
            trt_fail = trt_fail_reason(trt_entry)

            agent_dir = os.path.join(OUT_DIR, f"2_{pid}")
            metrics_path = os.path.join(agent_dir, "global_best_metrics_50.json")
            kernel_path = os.path.join(agent_dir, "global_best_kernel_50.py")
            metrics = None
            if os.path.isfile(metrics_path):
                try:
                    with open(metrics_path, "r") as f:
                        metrics = json.load(f)
                except Exception:
                    metrics = None
            elif not os.path.isdir(agent_dir):
                metrics = None

            agent_runtime = metrics.get("runtime") if metrics else None
            a_fail = ""
            if metrics is None:
                if not os.path.isdir(agent_dir):
                    a_fail = "Agent did not run"
                else:
                    a_fail = "Agent optimization failure"
            else:
                a_fail = agent_fail_reason(metrics)

            if trt_fail and a_fail == "Agent did not run":
                fail = trt_fail
            else:
                fail = "; ".join(x for x in (trt_fail, a_fail) if x)

            ratio = ""
            if trt_ok and agent_runtime and agent_runtime > 0 and not a_fail:
                ratio = round(trt_mean / agent_runtime, 4)

            impl = detect_impl(kernel_path) if metrics and not a_fail else ""

            agent_cell = round(agent_runtime, 4) if (agent_runtime and agent_runtime > 0) else ""

            rows.append((
                fn,
                round(trt_mean, 4) if trt_mean else "",
                agent_cell,
                ratio,
                fail,
                impl,
            ))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "level2"
    headers = ["Name", "trt_baseline (ms)", "Agent result", "ratio", "Fail", "Implementation"]
    ws.append(headers)
    bold = Font(bold=True)
    fill = PatternFill("solid", fgColor="DDDDDD")
    for col_idx in range(1, len(headers) + 1):
        c = ws.cell(row=1, column=col_idx)
        c.font = bold
        c.fill = fill
        c.alignment = Alignment(horizontal="center")

    def pid_of(name):
        m = re.match(r"^(\d+)_", name)
        return int(m.group(1)) if m else 10**9

    rows.sort(key=lambda r: pid_of(r[0]))
    for r in rows:
        ws.append(r)

    widths = [62, 20, 16, 10, 38, 18]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    ws.freeze_panes = "A2"
    wb.save(XLSX)
    print(f"Wrote {len(rows)} rows to {XLSX}")


if __name__ == "__main__":
    main()
